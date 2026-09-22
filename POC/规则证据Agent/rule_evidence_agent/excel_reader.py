from __future__ import annotations

from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook


STRUCTURED_REQUIRED_COLUMNS = {
    "规则编号",
    "规则名称",
    "医保项目编码",
    "项目名称",
    "负面清单规则原文",
    "生效日期",
    "规则版本",
    "发布状态",
}
EXPORT_REQUIRED_COLUMNS = {
    "住院门诊号",
    "就诊号",
    "mdtrt_id",
    "违规内容",
    "规则名称",
    "医保项目编码",
    "医保项目名称",
}
SIMPLE_REQUIRED_COLUMNS = {"住院门诊号", "就诊号", "违规内容"}


@dataclass(frozen=True)
class RuleRecord:
    rule_id: str
    source_rule_name: str
    item_code: str
    item_name: str
    rule_text: str
    effective_date: date | None
    source_version: str
    mdtrt_id: str = ""
    medins_id: str = ""
    visit_no: str = ""


def _as_identifier(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()

def _as_date(value: object) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    return None


def read_published_rules(path: Path) -> list[RuleRecord]:
    if not path.exists():
        raise FileNotFoundError(f"负面清单不存在: {path}")

    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook[workbook.sheetnames[0]]

    header_row = None
    headers: list[str] = []
    input_format = ""
    for row_index, row in enumerate(sheet.iter_rows(min_row=1, max_row=20, values_only=True), start=1):
        values = [str(value).strip() if value is not None else "" for value in row]
        value_set = set(values)
        if STRUCTURED_REQUIRED_COLUMNS.issubset(value_set):
            header_row = row_index
            headers = values
            input_format = "structured"
            break
        if EXPORT_REQUIRED_COLUMNS.issubset(value_set):
            header_row = row_index
            headers = values
            input_format = "hospital_export"
            break
        if SIMPLE_REQUIRED_COLUMNS.issubset(value_set):
            header_row = row_index
            headers = values
            input_format = "simple"
            break
    if header_row is None:
        raise ValueError(
            "前 20 行未找到可识别表头；支持完整规则表，"
            f"或三列表 {sorted(SIMPLE_REQUIRED_COLUMNS)}"
        )

    index = {name: pos for pos, name in enumerate(headers) if name}
    records: list[RuleRecord] = []
    for row in sheet.iter_rows(min_row=header_row + 1, values_only=True):
        if not any(value is not None for value in row):
            continue
        if input_format == "simple":
            rule_text = str(row[index["违规内容"]] or "").strip()
            if not rule_text:
                continue
            mdtrt_id = _as_identifier(row[index["就诊号"]]) if "就诊号" in index else ""
            if "mdtrt_id" in index:
                mdtrt_id = _as_identifier(row[index["mdtrt_id"]])
                if not mdtrt_id:
                    continue
            rule_id = f"RULE-{len(records) + 1:03d}"
            records.append(
                RuleRecord(
                    rule_id=rule_id,
                    source_rule_name=f"待拆解规则 {rule_id}",
                    item_code="",
                    item_name="",
                    rule_text=rule_text,
                    effective_date=None,
                    source_version="POC",
                    mdtrt_id=mdtrt_id,
                    visit_no=_as_identifier(row[index["住院门诊号"]]) if "住院门诊号" in index else "",
                )
            )
            continue

        if input_format == "hospital_export":
            rule_text = str(row[index["违规内容"]] or "").strip()
            if not rule_text:
                continue
            mdtrt_id = _as_identifier(row[index["mdtrt_id"]])
            if not mdtrt_id:
                continue
            rule_id = f"RULE-{len(records) + 1:03d}"
            item_name = str(row[index["医保项目名称"]] or "").strip()
            if not item_name and "医院项目名称" in index:
                item_name = str(row[index["医院项目名称"]] or "").strip()
            source_rule_name = str(row[index["规则名称"]] or "").strip() or f"待拆解规则 {rule_id}"
            records.append(
                RuleRecord(
                    rule_id=rule_id,
                    source_rule_name=source_rule_name,
                    item_code=str(row[index["医保项目编码"]] or "").strip(),
                    item_name=item_name,
                    rule_text=rule_text,
                    effective_date=_as_date(row[index["费用日期"]]) if "费用日期" in index else None,
                    source_version="南山人医-数据导出",
                    mdtrt_id=mdtrt_id,
                    medins_id=_as_identifier(row[index["medins_id"]]) if "medins_id" in index else "",
                    visit_no=_as_identifier(row[index["住院门诊号"]]) if "住院门诊号" in index else "",
                )
            )
            continue

        status = str(row[index["发布状态"]] or "").strip()
        if status != "已发布":
            continue
        record = RuleRecord(
            rule_id=str(row[index["规则编号"]] or "").strip(),
            source_rule_name=str(row[index["规则名称"]] or "").strip(),
            item_code=str(row[index["医保项目编码"]] or "").strip(),
            item_name=str(row[index["项目名称"]] or "").strip(),
            rule_text=str(row[index["负面清单规则原文"]] or "").strip(),
            effective_date=_as_date(row[index["生效日期"]]),
            source_version=str(row[index["规则版本"]] or "").strip(),
            mdtrt_id=_as_identifier(row[index["mdtrt_id"]]) if "mdtrt_id" in index else "",
            medins_id=_as_identifier(row[index["medins_id"]]) if "medins_id" in index else "",
            visit_no=_as_identifier(row[index["住院门诊号"]]) if "住院门诊号" in index else "",
        )
        missing = [
            name
            for name, value in {
                "规则编号": record.rule_id,
                "规则名称": record.source_rule_name,
                "医保项目编码": record.item_code,
                "项目名称": record.item_name,
                "负面清单规则原文": record.rule_text,
                "规则版本": record.source_version,
            }.items()
            if not value
        ]
        if missing:
            raise ValueError(f"规则行缺少必填值 {missing}: {record.rule_id or '<无编号>'}")
        records.append(record)

    if not records:
        raise ValueError("负面清单中没有可处理的规则")
    if len({record.rule_id for record in records}) != len(records):
        raise ValueError("负面清单存在重复规则编号")
    return records
