from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook


REQUIRED_COLUMNS = ("一级证据域", "二级证据标签")


@dataclass(frozen=True)
class EvidenceTag:
    primary_domain: str
    secondary_tag: str

    @property
    def pair(self) -> tuple[str, str]:
        return self.primary_domain, self.secondary_tag

    def to_prompt_dict(self) -> dict[str, str]:
        return {
            "一级证据域": self.primary_domain,
            "二级证据标签": self.secondary_tag,
        }


def load_tags(path: Path) -> list[EvidenceTag]:
    if not path.exists():
        raise FileNotFoundError(f"标签目录不存在: {path}")

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook[workbook.sheetnames[0]]
        sheet.reset_dimensions()
        rows = sheet.iter_rows(values_only=True)
        try:
            header = next(rows)
        except StopIteration as exc:
            raise ValueError("标签目录为空") from exc

        header_map = {
            str(value).strip(): index
            for index, value in enumerate(header)
            if value is not None
        }
        missing = [column for column in REQUIRED_COLUMNS if column not in header_map]
        if missing:
            raise ValueError(f"标签目录缺少列: {', '.join(missing)}")

        tags: list[EvidenceTag] = []
        seen: set[tuple[str, str]] = set()
        for row_number, row in enumerate(rows, start=2):
            primary_index = header_map["一级证据域"]
            secondary_index = header_map["二级证据标签"]
            primary = "" if primary_index >= len(row) or row[primary_index] is None else str(row[primary_index]).strip()
            secondary = "" if secondary_index >= len(row) or row[secondary_index] is None else str(row[secondary_index]).strip()
            if not primary and not secondary:
                continue
            if not primary or not secondary:
                raise ValueError(f"标签目录第 {row_number} 行存在必填字段为空")
            pair = (primary, secondary)
            if pair in seen:
                raise ValueError(f"标签目录存在重复标签: {primary}—{secondary}")
            seen.add(pair)
            tags.append(EvidenceTag(primary, secondary))

        if not tags:
            raise ValueError("标签目录中没有可用标签")
        return tags
    finally:
        workbook.close()
