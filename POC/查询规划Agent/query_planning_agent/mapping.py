from __future__ import annotations

import re
from pathlib import Path

from openpyxl import load_workbook


def load_expense_tags(path: Path) -> set[tuple[str, str]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        worksheet = workbook["拆分映射"]
        header: dict[str, int] | None = None
        tags: set[tuple[str, str]] = set()
        for values in worksheet.iter_rows(values_only=True):
            if values and values[0] == "mapping_id":
                header = {
                    str(value): index
                    for index, value in enumerate(values)
                    if value is not None
                }
                continue
            if header is None or not values or not values[0]:
                continue
            view_name = str(values[header["对应的视图"]] or "").strip()
            source_cells = [
                str(values[header[name]] or "")
                for name in ("主要来源", "次要来源")
                if name in header
            ]
            if view_name != "费用明细" and not any(
                re.search(r"DS-S-002\s*费用明细", cell) for cell in source_cells
            ):
                continue
            tags.add(
                (
                    str(values[header["一级标签名"]] or "").strip(),
                    str(values[header["二级标签名"]] or "").strip(),
                )
            )
        return tags
    finally:
        workbook.close()

