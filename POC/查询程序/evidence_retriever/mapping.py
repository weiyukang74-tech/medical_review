from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
import re

from openpyxl import load_workbook


@dataclass(frozen=True)
class SourceRoute:
    mapping_id: str
    medins_id: str
    mdtrt_id: str
    primary_domain: str
    secondary_tag: str
    source_level: str
    source_code: str
    source_name: str
    view_name: str
    document_kind: str
    record_name: str | None
    record_type: str | None
    mapping_status: str


def _clean_optional(value: object) -> str | None:
    text = str(value or "").strip()
    if not text or text in {"—", "不适用", "未匹配"}:
        return None
    return text


def _parse_source(value: object) -> tuple[str, str]:
    text = str(value or "").strip()
    if not text or text == "—":
        raise ValueError("映射表来源为空")
    match = re.match(r"^(DS-(?:S|M-(?:IP|OP))-\d{3})\s*(.+)$", text)
    if not match:
        raise ValueError(f"无法解析来源编码和名称：{text}")
    return match.group(1), match.group(2).strip()


def load_routes(path: Path) -> list[SourceRoute]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    if "拆分映射" not in workbook.sheetnames:
        raise ValueError("映射表缺少“拆分映射”工作表")
    worksheet = workbook["拆分映射"]
    header = None
    routes: list[SourceRoute] = []
    for values in worksheet.iter_rows(values_only=True):
        if values and values[0] == "mapping_id":
            header = {str(value): index for index, value in enumerate(values) if value is not None}
            continue
        if header is None or not values or not values[0]:
            continue
        raw_level = str(values[header["来源级别"]]).strip()
        if raw_level in {"PRIMARY", "主要来源"}:
            level = "PRIMARY"
        elif raw_level in {"SECONDARY", "次要来源"}:
            level = "SECONDARY"
        else:
            raise ValueError(f"无法识别来源级别：{raw_level}")
        source_cell = values[header["主要来源"]] if level == "PRIMARY" else values[header["次要来源"]]
        source_code, source_name = _parse_source(source_cell)
        routes.append(
            SourceRoute(
                mapping_id=str(values[header["mapping_id"]]).strip(),
                medins_id=str(values[header["medinsId"]]).strip(),
                mdtrt_id=str(values[header["mdtrtId"]]).strip(),
                primary_domain=str(values[header["一级标签名"]]).strip(),
                secondary_tag=str(values[header["二级标签名"]]).strip(),
                source_level=level,
                source_code=source_code,
                source_name=source_name,
                view_name=str(values[header["对应的视图"]]).strip(),
                document_kind=str(values[header["文书类型"]]).strip(),
                record_name=_clean_optional(values[header["匹配到的recordName"]]),
                record_type=_clean_optional(values[header["recordType"]]),
                mapping_status=str(values[header["匹配状态"]]).strip(),
            )
        )
    if not routes:
        raise ValueError("映射表中没有可用的映射记录")
    return routes


def index_routes(routes: list[SourceRoute]) -> dict[tuple[str, str], dict[str, list[SourceRoute]]]:
    result: dict[tuple[str, str], dict[str, list[SourceRoute]]] = {}
    for route in routes:
        levels = result.setdefault((route.primary_domain, route.secondary_tag), {"PRIMARY": [], "SECONDARY": []})
        levels[route.source_level].append(route)
    return result
