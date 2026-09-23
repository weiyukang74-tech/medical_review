from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import load_workbook


DEFAULT_FIELD_DICTIONARY = (
    Path(__file__).resolve().parents[3] / "datat" / "标准病历视图.xlsx"
)


class StandardViewFieldDictionary:
    def __init__(self, path: Path | None = None):
        self.path = path or DEFAULT_FIELD_DICTIONARY
        self.fields: dict[str, dict[str, str]] = {}
        if self.path.exists():
            self._load()

    @staticmethod
    def _normalize(field_name: object) -> str:
        return "".join(
            character.casefold()
            for character in str(field_name or "")
            if character.isalnum()
        )

    def _load(self) -> None:
        workbook = load_workbook(self.path, read_only=True, data_only=True)
        try:
            worksheet = workbook.active
            current_view = ""
            for row in worksheet.iter_rows(min_row=2, values_only=True):
                view_name = str(row[0] or "").strip()
                field_name = str(row[1] or "").strip()
                description = str(row[2] or "").strip()
                if view_name:
                    current_view = view_name
                if not field_name:
                    continue
                self.fields[self._normalize(field_name)] = {
                    "field_name": field_name,
                    "description": description,
                    "view_name": current_view,
                }
        finally:
            workbook.close()

    def describe(self, field_name: str) -> dict[str, str]:
        entry = self.fields.get(self._normalize(field_name))
        if entry:
            return entry
        return {
            "field_name": field_name,
            "description": field_name,
            "view_name": "",
        }

    def labels_for(self, field_names: list[str]) -> dict[str, str]:
        return {
            field_name: self.describe(field_name)["description"]
            for field_name in field_names
        }
