from __future__ import annotations

import json
import re
import xml.etree.ElementTree as ET
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
import warnings

from openpyxl import load_workbook


SECTION_TO_VIEW = {
    "患者信息": "患者信息",
    "费用明细": "费用明细",
    "手术操作": "手术操作",
    "诊断信息": "诊断信息",
    "检验报告明细": "检验报告明细",
    "病历信息": "病历信息",
}


def _local_name(tag: str) -> str:
    return tag.rsplit("}", 1)[-1]


def _parse_value(value: str | None) -> object:
    if value is None:
        return None
    text = value.strip()
    if not text:
        return ""
    if text.startswith("{") or text.startswith("["):
        try:
            return json.loads(text)
        except json.JSONDecodeError:
            return text
    return text


@dataclass(frozen=True)
class MedicalDataset:
    section_name: str
    records: tuple[dict[str, object], ...]


class MarkdownMedicalData:
    def __init__(self, path: Path):
        self.path = path
        self.is_xlsx_directory = path.is_dir()
        self.datasets = {} if self.is_xlsx_directory else self._load(path)
        self._xlsx_paths = (
            [
                item
                for item in sorted(path.glob("*.xlsx"))
                if not item.name.startswith("~$") and "负面清单" not in item.stem
            ]
            if self.is_xlsx_directory
            else []
        )
        self._view_paths: dict[str, list[Path]] | None = None
        self._xlsx_query_cache: dict[tuple[str, str, str], list[dict[str, object]]] = {}

    def _index_xlsx_views(self) -> dict[str, list[Path]]:
        if self._view_paths is not None:
            return self._view_paths
        view_paths: dict[str, list[Path]] = {}
        for path in self._xlsx_paths:
            with warnings.catch_warnings():
                warnings.simplefilter("ignore", UserWarning)
                workbook = load_workbook(path, read_only=True, data_only=True)
            try:
                for worksheet in workbook.worksheets:
                    header = _read_header(worksheet)
                    view_name = _detect_view_name(path, header, worksheet.title)
                    if view_name:
                        view_paths.setdefault(view_name, []).append(path)
            finally:
                workbook.close()
        self._view_paths = view_paths
        return view_paths

    def _query_xlsx_view(
        self,
        view_name: str,
        medins_id: str,
        mdtrt_id: str,
    ) -> list[dict[str, object]]:
        cache_key = (view_name, medins_id, mdtrt_id)
        if cache_key in self._xlsx_query_cache:
            return self._xlsx_query_cache[cache_key]
        records: list[dict[str, object]] = []
        for path in self._index_xlsx_views().get(view_name, []):
            with warnings.catch_warnings():
                warnings.simplefilter("ignore", UserWarning)
                workbook = load_workbook(path, read_only=True, data_only=True)
            try:
                for worksheet in workbook.worksheets:
                    header = _read_header(worksheet)
                    if not header or "mdtrtId" not in header:
                        continue
                    medins_index = header.index("medinsId") if "medinsId" in header else None
                    mdtrt_index = header.index("mdtrtId")
                    for raw_row in worksheet.iter_rows(values_only=True):
                        values = [_coerce_cell(value) for value in raw_row]
                        if len(values) <= mdtrt_index or str(values[mdtrt_index] or "") != mdtrt_id:
                            continue
                        if medins_index is not None and str(values[medins_index] or "") != medins_id:
                            continue
                        records.append(_row_to_record(header, values))
            finally:
                workbook.close()
        self._xlsx_query_cache[cache_key] = records
        return records

    @classmethod
    def _load_xlsx_directory(cls, directory: Path) -> dict[str, MedicalDataset]:
        datasets: dict[str, MedicalDataset] = {}
        for path in sorted(directory.glob("*.xlsx")):
            if path.name.startswith("~$") or "负面清单" in path.stem:
                continue
            workbook = load_workbook(path, read_only=True, data_only=True)
            for worksheet in workbook.worksheets:
                rows = worksheet.iter_rows(values_only=True)
                header: list[str] | None = None
                records: list[dict[str, object]] = []
                for raw_row in rows:
                    values = [_coerce_cell(value) for value in raw_row]
                    if header is None:
                        candidate = [str(value).strip() if value is not None else "" for value in values]
                        if not any(candidate):
                            continue
                        header = candidate
                        continue
                    record = {
                        key: _parse_value(value) if key == "recordContent" and isinstance(value, str) else value
                        for key, value in zip(header, values)
                        if key
                    }
                    if any(value not in (None, "") for value in record.values()):
                        records.append(record)
                view_name = _detect_view_name(path, header or [], worksheet.title)
                if view_name is None:
                    continue
                current = datasets.get(view_name)
                merged_records = list(current.records) if current else []
                merged_records.extend(records)
                datasets[view_name] = MedicalDataset(
                    section_name=worksheet.title,
                    records=tuple(merged_records),
                )
            workbook.close()
        return datasets

    @staticmethod
    def _load(path: Path) -> dict[str, MedicalDataset]:
        text = path.read_text(encoding="utf-8")
        matches = list(re.finditer(r"(?m)^###\s+(.+?)\s*$", text))
        datasets: dict[str, MedicalDataset] = {}
        for index, match in enumerate(matches):
            title = match.group(1).strip()
            body_start = match.end()
            body_end = matches[index + 1].start() if index + 1 < len(matches) else len(text)
            body = text[body_start:body_end]
            code_match = re.search(r"```(?:xml)?\s*(.*?)```", body, re.DOTALL)
            if not code_match:
                continue
            xml_text = code_match.group(1).strip()
            if not xml_text.startswith("<?xml"):
                continue
            root = ET.fromstring(xml_text)
            records: list[dict[str, object]] = []
            for element in root.iter():
                if _local_name(element.tag) != "data":
                    continue
                record: dict[str, object] = {}
                for child in list(element):
                    record[_local_name(child.tag)] = _parse_value(child.text)
                records.append(record)
            normalized = title.split("_", 1)[0].strip()
            view_name = SECTION_TO_VIEW.get(normalized, normalized)
            datasets[view_name] = MedicalDataset(title, tuple(records))
        return datasets

    def query_structured(self, view_name: str, medins_id: str, mdtrt_id: str) -> list[dict[str, object]]:
        if self.is_xlsx_directory:
            return self._query_xlsx_view(view_name, medins_id, mdtrt_id)
        dataset = self.datasets.get(view_name)
        if dataset is None:
            return []
        return [
            record
            for record in dataset.records
            if str(record.get("medinsId", "")) == medins_id
            and str(record.get("mdtrtId", "")) == mdtrt_id
        ]

    def resolve_medins_id(self, mdtrt_id: str) -> str | None:
        if self.is_xlsx_directory:
            for path in self._index_xlsx_views().get("患者信息", []):
                with warnings.catch_warnings():
                    warnings.simplefilter("ignore", UserWarning)
                    workbook = load_workbook(path, read_only=True, data_only=True)
                try:
                    for worksheet in workbook.worksheets:
                        header = _read_header(worksheet)
                        if "mdtrtId" not in header or "medinsId" not in header:
                            continue
                        mdtrt_index = header.index("mdtrtId")
                        medins_index = header.index("medinsId")
                        for raw_row in worksheet.iter_rows(values_only=True):
                            values = [_coerce_cell(value) for value in raw_row]
                            if (
                                len(values) > max(mdtrt_index, medins_index)
                                and str(values[mdtrt_index] or "") == mdtrt_id
                            ):
                                medins_id = str(values[medins_index] or "").strip()
                                if medins_id:
                                    return medins_id
                finally:
                    workbook.close()
            return None

        for dataset in self.datasets.values():
            for record in dataset.records:
                if str(record.get("mdtrtId", "")) == mdtrt_id:
                    medins_id = str(record.get("medinsId", "")).strip()
                    if medins_id:
                        return medins_id
        return None

    def query_records(
        self,
        medins_id: str,
        mdtrt_id: str,
        record_name: str | None,
        record_type: str | None,
    ) -> list[dict[str, object]]:
        if self.is_xlsx_directory:
            records = self._query_xlsx_view("病历信息", medins_id, mdtrt_id)
            return [
                record
                for record in records
                if (not record_name or str(record.get("recordName", "")) == record_name)
                and (not record_type or str(record.get("recordType", "")) == record_type)
            ]
        dataset = self.datasets.get("病历信息")
        if dataset is None:
            return []
        result: list[dict[str, object]] = []
        for record in dataset.records:
            if str(record.get("medinsId", "")) != medins_id or str(record.get("mdtrtId", "")) != mdtrt_id:
                continue
            if record_name and str(record.get("recordName", "")) != record_name:
                continue
            if record_type and str(record.get("recordType", "")) != record_type:
                continue
            result.append(record)
        return result


def _coerce_cell(value: object) -> object:
    if isinstance(value, (datetime, date)):
        return value.isoformat(sep=" ") if isinstance(value, datetime) else value.isoformat()
    return value


def _read_header(worksheet: object) -> list[str]:
    for raw_row in worksheet.iter_rows(values_only=True):
        values = [_coerce_cell(value) for value in raw_row]
        header = [str(value).strip() if value is not None else "" for value in values]
        if any(header):
            return header
    return []


def _row_to_record(header: list[str], values: list[object]) -> dict[str, object]:
    return {
        key: _parse_value(value) if key == "recordContent" and isinstance(value, str) else value
        for key, value in zip(header, values)
        if key
    }


def _detect_view_name(path: Path, headers: list[str], sheet_name: str) -> str | None:
    header_set = set(headers)
    if "recordContent" in header_set:
        return "病历信息"
    if "diseName" in header_set:
        return "诊断信息"
    if {"patnName", "gend"}.issubset(header_set):
        return "患者信息"
    if {"rxId", "mdtrtId"}.issubset(header_set) or "hilistName" in header_set:
        return "费用明细"

    filename = path.stem
    if "检验报告" in filename or "检验报告" in sheet_name:
        return "检验报告明细"
    if "手术操作" in filename or "手术操作" in sheet_name:
        return "手术操作"
    if "诊断信息" in filename or "诊断信息" in sheet_name:
        return "诊断信息"
    if "患者信息" in filename or "患者信息" in sheet_name:
        return "患者信息"
    if "费用明细" in filename or "费用明细" in sheet_name:
        return "费用明细"
    if "病历信息" in filename or "病例信息" in filename or "病历信息" in sheet_name:
        return "病历信息"
    return None
